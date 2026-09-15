import base64
from datetime import date
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from PIL import Image
import pytest
from sqlalchemy import text
from sqlmodel import Session, select

import test_local_staff_auth as staff_tests
import test_calendar_feature as calendar_tests
from calendar_import import parse_ics, router as import_router
from local_auth import activate_staff_password
from meeting_note_export import export_note
from models import Calendar, CalendarImportBatch, CalendarImportSource, Event, User


@pytest.fixture
def staff():
    case = staff_tests.LocalStaffAuthenticationTests()
    case.setUp()
    case._activate()
    case._login()
    try:
        yield case
    finally:
        case.tearDown()


@pytest.mark.parametrize("order", [199, 200, 220, 240, 1000])
def test_staff_activation_and_login_do_not_depend_on_display_order(staff, order):
    with Session(staff.engine) as session:
        user = User(email=f"staff-{order}@example.test", display_name="新規職員", staff_sort_order=order)
        session.add(user)
        session.commit()
        user_id = user.id
    response = staff.client.post(f"/staff/users/{user_id}/authentication/activate", data={"login_id": "new-staff", "reason": "採用"})
    assert response.status_code == 200
    code = staff._action_code_from(response)
    with Session(staff.engine) as session:
        activate_staff_password(session, activation_code=code, password=staff.PASSWORD, password_confirmation=staff.PASSWORD)
    assert staff.client.post("/staff/login", data={"login_id": "new-staff", "password": staff.PASSWORD}, follow_redirects=False).status_code == 303
    assert staff.client.get("/protected").status_code == 200


def test_display_order_change_keeps_session_but_deactivation_revokes_it(staff):
    with Session(staff.engine) as session:
        user = session.get(User, staff.user_id)
        user.staff_sort_order = 500
        session.add(user)
        session.commit()
    assert staff.client.get("/protected").status_code == 200
    with Session(staff.engine) as session:
        user = session.get(User, staff.user_id)
        user.is_active = False
        session.add(user)
        session.commit()
    assert staff.client.get("/protected", follow_redirects=False).status_code == 401


def ics(*events):
    return ("BEGIN:VCALENDAR\r\nVERSION:2.0\r\nX-WR-TIMEZONE:Asia/Tokyo\r\n" + "".join("BEGIN:VEVENT\r\n" + item + "\r\nEND:VEVENT\r\n" for item in events) + "END:VCALENDAR\r\n").encode()


SINGLE = "UID:single\r\nDTSTART;VALUE=DATE:20260915\r\nDTEND;VALUE=DATE:20260917\r\nSUMMARY:宿泊行事\r\nCLASS:PRIVATE"
REPEAT = "UID:repeat\r\nDTSTART;TZID=Asia/Tokyo:20260914T090000\r\nDTEND;TZID=Asia/Tokyo:20260914T100000\r\nRRULE:FREQ=DAILY;COUNT=4\r\nEXDATE;TZID=Asia/Tokyo:20260915T090000\r\nSUMMARY:朝会"
MOVED = "UID:repeat\r\nRECURRENCE-ID;TZID=Asia/Tokyo:20260916T090000\r\nDTSTART;TZID=Asia/Tokyo:20260916T110000\r\nDTEND;TZID=Asia/Tokyo:20260916T120000\r\nSUMMARY:時刻変更"


def test_ics_preserves_all_day_exclusions_and_moved_occurrences():
    rows = parse_ics(ics(SINGLE, REPEAT, MOVED), date(2026, 9, 14), date(2026, 9, 17))
    assert len(rows) == 4
    overnight = next(row for row in rows if row["title"] == "宿泊行事")
    assert overnight["is_all_day"] and overnight["visibility"] == "private"
    assert overnight["start_at"] == "2026-09-14T15:00:00+00:00"
    assert overnight["end_at"] == "2026-09-16T15:00:00+00:00"
    moved = next(row for row in rows if row["title"] == "時刻変更")
    assert moved["start_at"] == "2026-09-16T02:00:00+00:00"
    original = parse_ics(ics(REPEAT), date(2026, 9, 16), date(2026, 9, 16))[0]
    assert moved["source_key"] == original["source_key"]


def test_ics_dst_month_end_and_private_events():
    content = ics("UID:dst\r\nDTSTART;TZID=America/New_York:20261031T090000\r\nDTEND;TZID=America/New_York:20261031T100000\r\nRRULE:FREQ=DAILY;COUNT=3\r\nSUMMARY:DST")
    rows = parse_ics(content, date(2026, 10, 31), date(2026, 11, 3))
    assert [r["start_at"][11:16] for r in rows] == ["13:00", "14:00", "14:00"]
    rows = parse_ics(ics("UID:month\r\nDTSTART:20260131T090000Z\r\nRRULE:FREQ=MONTHLY;COUNT=3\r\nSUMMARY:月末"), date(2026, 1, 1), date(2026, 6, 1))
    assert [r["start_at"][:10] for r in rows] == ["2026-01-31", "2026-03-31", "2026-05-31"]


def test_single_event_identity_survives_a_date_change():
    original = parse_ics(ics(SINGLE), date(2026, 9, 1), date(2026, 9, 30))[0]
    changed = parse_ics(ics(SINGLE.replace("20260915", "20260916")), date(2026, 9, 1), date(2026, 9, 30))[0]
    assert original["source_key"] == changed["source_key"]


@pytest.mark.parametrize("event", [SINGLE.replace("UID:single\r\n", ""), REPEAT.replace("FREQ=DAILY", "FREQ=SECONDLY"), MOVED])
def test_ics_rejects_unsupported_or_incomplete_data(event):
    with pytest.raises(ValueError):
        parse_ics(ics(event), date(2026, 9, 1), date(2026, 10, 1))


@pytest.fixture
def calendar():
    case = calendar_tests.CalendarFeatureTests()
    case.setUp()
    case.app.include_router(import_router)
    case._login(case.user_a_id)
    try:
        yield case
    finally:
        case.tearDown()


def preview(calendar):
    response = calendar.client.post("/calendar/import/preview", data={"calendar_id": str(calendar.a_personal_id), "start": "2026-09-01", "end": "2026-09-30"}, files={"file": ("calendar.ics", ics(SINGLE, REPEAT, MOVED))})
    assert response.status_code == 200
    return response.context["batch"].id


def test_calendar_preview_commit_replay_and_import_permissions(calendar):
    batch_id = preview(calendar)
    with Session(calendar.engine) as session:
        before = len(session.exec(select(Event)).all())
    response = calendar.client.post("/calendar/import/commit", data={"batch_id": str(batch_id)}, follow_redirects=False)
    assert response.status_code == 303
    with Session(calendar.engine) as session:
        assert len(session.exec(select(Event)).all()) == before + 4
        assert len(session.exec(select(CalendarImportSource)).all()) == 4
    assert "取込済み" in calendar.client.post("/calendar/import/commit", data={"batch_id": str(batch_id)}).text
    repeated = preview(calendar)
    assert calendar.client.post("/calendar/import/commit", data={"batch_id": str(repeated)}, follow_redirects=False).status_code == 303
    with Session(calendar.engine) as session:
        assert len(session.exec(select(Event)).all()) == before + 4
    owned = preview(calendar)
    calendar._login(calendar.user_b_id)
    assert "取込済み" in calendar.client.post("/calendar/import/commit", data={"batch_id": str(owned)}).text
    with Session(calendar.engine) as session:
        assert session.get(CalendarImportBatch, owned).used_at is None
    response = calendar.client.post("/calendar/import/preview", data={"calendar_id": str(calendar.a_personal_id), "start": "2026-09-01", "end": "2026-09-30"}, files={"file": ("calendar.ics", ics(SINGLE))})
    assert response.status_code == 403


def test_deleting_an_imported_calendar_clears_import_references(calendar):
    batch_id = preview(calendar)
    assert calendar.client.post("/calendar/import/commit", data={"batch_id": str(batch_id)}, follow_redirects=False).status_code == 303
    with calendar.engine.connect() as connection:
        connection.execute(text("PRAGMA foreign_keys=ON"))
    response = calendar.client.post(f"/calendars/{calendar.a_personal_id}/delete")
    assert response.status_code == 200
    with Session(calendar.engine) as session:
        assert session.get(Calendar, calendar.a_personal_id) is None
        assert not session.exec(select(CalendarImportSource)).all()
        assert not session.exec(select(CalendarImportBatch)).all()
        assert not session.execute(text("PRAGMA foreign_key_check")).all()


def test_docx_and_markdown_preserve_headings_text_images_and_links():
    image = BytesIO()
    Image.new("RGB", (40, 30), "orange").save(image, "PNG")
    data_url = "data:image/png;base64," + base64.b64encode(image.getvalue()).decode()
    ops = [{"insert": "議題"}, {"insert": "\n", "attributes": {"header": 1}}, {"insert": "確認事項", "attributes": {"bold": True}}, {"insert": "\n", "attributes": {"list": "bullet"}}, {"insert": "参考", "attributes": {"link": "https://example.test/a"}}, {"insert": "\n"}, {"insert": {"image": data_url}}, {"insert": "\n"}]
    docx, mime, extension = export_note("職員会議", ops, "docx")
    document = Document(BytesIO(docx))
    assert document.paragraphs[0].text == "職員会議"
    assert document.paragraphs[1].style.name == "Heading 1"
    assert document.paragraphs[2].runs[0].bold
    assert len(document.inline_shapes) == 1 and extension == "docx"
    with ZipFile(BytesIO(docx)) as archive:
        assert b"https://example.test/a" in archive.read("word/_rels/document.xml.rels")
    md, mime, extension = export_note("職員会議", ops, "md")
    assert extension == "zip"
    with ZipFile(BytesIO(md)) as archive:
        text = archive.read("meeting-note.md").decode()
        assert "- **確認事項**" in text
        assert "[参考](https://example.test/a)" in text
        assert "images/image-1.png" in archive.namelist()


def test_export_rejects_invalid_images_and_does_not_emit_active_links():
    with pytest.raises(ValueError):
        export_note("議事録", [{"insert": {"image": "file:///secret"}}], "docx")
    content, _, extension = export_note("議事録", [{"insert": "<script>本文</script>", "attributes": {"link": "javascript:alert(1)"}}], "md")
    assert extension == "md" and b"javascript:" not in content and b"<script>" not in content
