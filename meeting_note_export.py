"""Convert the visible Quill document to DOCX or Markdown without changing the note."""
import base64
from io import BytesIO
import re
from urllib.parse import quote, urlsplit
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE
from PIL import Image


def safe_url(value):
    value = str(value or "")
    return value if urlsplit(value).scheme.lower() in {"https", "http", "mailto"} else ""


def blocks_from_ops(ops):
    blocks, parts = [], []
    characters = 0
    for op in ops:
        value, attrs = op.get("insert"), op.get("attributes") or {}
        if not isinstance(attrs, dict):
            raise ValueError("書式データが不正です。")
        if isinstance(value, str):
            characters += len(value)
            if characters > 500000 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value):
                raise ValueError("本文が長すぎるか、出力できない制御文字が含まれています。")
            chunks = value.split("\n")
            for index, chunk in enumerate(chunks):
                if chunk:
                    parts.append((chunk, attrs))
                if index < len(chunks) - 1:
                    blocks.append((parts, attrs))
                    parts = []
        elif isinstance(value, dict) and set(value) == {"image"}:
            parts.append((value, attrs))
        else:
            raise ValueError("出力できない埋め込みが含まれています。画像と文章に対応しています。")
    if parts:
        blocks.append((parts, {}))
    return blocks


def image_data(value):
    if not isinstance(value, str) or len(value) > 15 * 1024 * 1024:
        raise ValueError("画像が大きすぎます。")
    match = re.fullmatch(r"data:image/(?:png|jpeg|jpg|webp|gif);base64,([A-Za-z0-9+/=\r\n]+)", value)
    if not match:
        if safe_url(value):
            return None
        raise ValueError("画像の形式を読み取れません。")
    try:
        raw = base64.b64decode(match.group(1), validate=False)
        with Image.open(BytesIO(raw)) as image:
            if image.width * image.height > 25_000_000:
                raise ValueError("画像の解像度が大きすぎます。")
            image.thumbnail((1600, 1600))
            normalized = BytesIO()
            image.convert("RGB").save(normalized, "PNG")
            return normalized.getvalue()
    except Exception as exc:
        raise ValueError("画像を読み取れません。") from exc


def markdown_escape(value):
    return re.sub(r"([\\`*_{}\[\]<>#!|])", r"\\\1", value)


def export_note(title, ops, format_name):
    blocks = blocks_from_ops(ops)
    if len(blocks) > 10000:
        raise ValueError("段落数が多すぎます。")
    assets = {}
    image_cache = {}

    def picture(value):
        if value not in image_cache:
            data = image_data(value)
            if data:
                if len(assets) >= 100 or sum(map(len, assets.values())) + len(data) > 20 * 1024 * 1024:
                    raise ValueError("画像は合計20MB・100枚以内で出力してください。")
                name = f"images/image-{len(assets) + 1}.png"
                assets[name] = data
                image_cache[value] = (name, data)
            else:
                image_cache[value] = (safe_url(value), None)
        return image_cache[value]

    if format_name == "md":
        lines = ["# " + markdown_escape(title), ""]
        for parts, attrs in blocks:
            text = ""
            for value, run_attrs in parts:
                if isinstance(value, dict):
                    name, _ = picture(value["image"])
                    text += f"![画像]({quote(name, safe='/:@')})"
                    continue
                value = markdown_escape(value)
                if run_attrs.get("code"):
                    value = "`" + value.replace("`", "\\`") + "`"
                if run_attrs.get("bold"):
                    value = "**" + value + "**"
                if run_attrs.get("italic"):
                    value = "*" + value + "*"
                if run_attrs.get("strike"):
                    value = "~~" + value + "~~"
                if link := safe_url(run_attrs.get("link")):
                    value = f"[{value}]({quote(link, safe='/:@?=&%')})"
                text += value
            heading = attrs.get("header")
            indent = min(max(int(attrs.get("indent", 0)), 0), 8)
            if heading in (1, 2, 3, 4, 5, 6):
                text = "#" * heading + " " + text
            elif attrs.get("list"):
                text = "  " * indent + ("1. " if attrs["list"] == "ordered" else "- ") + text
            elif attrs.get("blockquote"):
                text = "> " + text
            elif attrs.get("code-block"):
                text = "    " + text
            lines.extend([text, ""])
        content = ("\n".join(lines).rstrip() + "\n").encode("utf-8")
        if not assets:
            return content, "text/markdown; charset=utf-8", "md"
        output = BytesIO()
        with ZipFile(output, "w", ZIP_DEFLATED) as archive:
            archive.writestr("meeting-note.md", content)
            for name, data in assets.items():
                archive.writestr(name, data)
        return output.getvalue(), "application/zip", "zip"

    if format_name != "docx":
        raise ValueError("出力形式を確認してください。")
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(22)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5", "Heading 6"):
        style = document.styles[name]
        style.font.name = "Yu Gothic"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Normal"].paragraph_format.line_spacing = 1.25
    document.add_heading(title, level=0)
    for parts, attrs in blocks:
        heading = attrs.get("header")
        style = f"Heading {heading}" if heading in (1, 2, 3, 4, 5, 6) else (
            "List Number" if attrs.get("list") == "ordered" else "List Bullet" if attrs.get("list") else "Normal")
        paragraph = document.add_paragraph(style=style)
        paragraph.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}.get(attrs.get("align"), WD_ALIGN_PARAGRAPH.LEFT)
        indent = min(max(int(attrs.get("indent", 0)), 0), 8)
        if indent:
            paragraph.paragraph_format.left_indent = Mm(6 * indent)
        for value, run_attrs in parts:
            if isinstance(value, dict):
                name, data = picture(value["image"])
                if data:
                    width = str(run_attrs.get("width", "100%"))
                    fraction = int(width[:-1]) / 100 if re.fullmatch(r"(?:25|50|75|100)%", width) else 1
                    with Image.open(BytesIO(data)) as image:
                        width_inches = min(6.1 * fraction, 8.0 * image.width / image.height)
                    paragraph.add_run().add_picture(BytesIO(data), width=Inches(width_inches))
                else:
                    paragraph.add_run("画像: " + name)
                continue
            run = paragraph.add_run(value)
            run.bold, run.italic = bool(run_attrs.get("bold")), bool(run_attrs.get("italic"))
            run.underline, run.font.strike = bool(run_attrs.get("underline")), bool(run_attrs.get("strike"))
            if run_attrs.get("code") or attrs.get("code-block"):
                run.font.name = "Consolas"
            if link := safe_url(run_attrs.get("link")):
                element = OxmlElement("w:hyperlink")
                element.set(qn("r:id"), paragraph.part.relate_to(link, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
                element.append(run._r)
                paragraph._p.append(element)
                run.underline = True
    document.core_properties.author = ""
    document.core_properties.title = title
    output = BytesIO()
    document.save(output)
    return output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"
